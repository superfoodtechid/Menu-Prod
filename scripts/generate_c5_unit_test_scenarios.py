import os
import sys
import copy
import openpyxl
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
INPUT_COMBINED = BASE_DIR / "data" / "exports" / "combined" / "es degan jeruk nipis dan 2 lainnya" / "C5. 2026-10-15 13:34 Unit Tes.xlsx"
OUTPUT_FILE = BASE_DIR / "C5. Unit Tes All Scenarios.xlsx"

# Known category IDs for target categories
CATEGORY_ID_MAP = {
    # GoFood
    ("GoFood", "Salad & Sehat"): "914c0d71-0bdd-4cbf-b315-74528b4abfc0",
    # GrabFood
    ("GrabFood", "Menu Unit Test Batch 2 Grab"): "IDCAT2026101500153822228",
    # ShopeeFood
    ("ShopeeFood", "Category B"): "3148481062563328",
    ("ShopeeFood", "Category A"): "3148480950140416",
}

def generate():
    if not INPUT_COMBINED.exists():
        print(f"Error: {INPUT_COMBINED} does not exist!")
        sys.exit(1)

    print(f"Loading input file: {INPUT_COMBINED}")
    wb = openpyxl.load_workbook(INPUT_COMBINED, data_only=False)
    ws = wb["Item"]

    # Column indices (1-indexed)
    col_ofd = 1
    col_outlet = 2
    col_sid = 5
    col_cat_id = 6
    col_cat = 7
    col_item_id = 8
    col_item = 9
    col_photo = 10
    col_desc = 11
    col_item_name_imp = 13
    col_avail = 18
    col_offline_price = 23
    col_offline_adj = 24
    col_curr_fake = 27
    col_new_fake = 40
    col_notes = 42

    # Track rows to delete (by item ID)
    items_to_delete = {
        "IDITE2026101405570200897",             # Grab: Nasi Goreng Terasi
        "3240705068845056",                      # Shopee: Ayam Geprek Iblis
        "82940cc5-cac5-4afb-9d3d-11f1cdeb21b0",  # GoFood: Dimsum Ayam Jamur Batch 2
    }

    max_r = ws.max_row
    max_c = ws.max_column

    rows_data = []
    for r in range(2, max_r + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, max_c + 1)]
        iid = str(row_vals[col_item_id - 1] or "").strip()
        iname = str(row_vals[col_item - 1] or "").strip()
        ofd = str(row_vals[col_ofd - 1] or "").strip()

        # Check if this row is targeted for deletion scenario
        if iid in items_to_delete:
            print(f"Deleting row {r} for DELETE_ITEM scenario: [{ofd}] {iname} (ID: {iid})")
            continue

        # ── GRABFOOD MUTATIONS (SID: 6-C7MAC3LDBFCCRE) ──
        if ofd == "GrabFood":
            # 1. Item A: CONTROL RULE: Price MUST be Rp 32.000 + OFFLINE_PRICE
            if "IDITE2026062704035134696" in iid:
                row_vals[col_new_fake - 1] = 32000
                row_vals[col_curr_fake - 1] = 32000
                row_vals[col_offline_price - 1] = 28000
                row_vals[col_offline_adj - 1] = 4000
                row_vals[col_notes - 1] = "[CONTROL: ITEM A RP 32K] [TEST: OFFLINE_PRICE]"
                print("Applied Grab Item A: Price set to Rp 32,000 + Offline Price Rp 28,000")

            # 2. Item B: PRICE_CHANGE (Normal <= 15%)
            elif "IDITE2026062704032966166" in iid:
                row_vals[col_new_fake - 1] = 35000
                row_vals[col_notes - 1] = "[TEST: PRICE_CHANGE NORMAL]"
                print("Applied Grab Item B: PRICE_CHANGE to Rp 35,000")

            # 3. Item C: PRICE_WARNING_STEP_PUSH (> 15%)
            elif "IDITE2026062704030116830" in iid:
                row_vals[col_new_fake - 1] = 45000
                row_vals[col_notes - 1] = "[TEST: PRICE_WARNING_STEP_PUSH]"
                print("Applied Grab Item C: PRICE_WARNING_STEP_PUSH to Rp 45,000")

            # 4. Item D: CATEGORY_CHANGE & NAME_CHANGE
            elif "IDITE2026062704023541089" in iid:
                target_cat = "Menu Unit Test Batch 2 Grab"
                row_vals[col_cat - 1] = target_cat
                row_vals[col_cat_id - 1] = CATEGORY_ID_MAP[("GrabFood", target_cat)]
                row_vals[col_item_name_imp - 1] = "Item D Spesial Grab"
                row_vals[col_notes - 1] = "[TEST: CATEGORY_CHANGE & NAME_CHANGE]"
                print(f"Applied Grab Item D: CATEGORY_CHANGE to '{target_cat}' (ID: {row_vals[col_cat_id - 1]}) & NAME_CHANGE")

        # ── SHOPEEFOOD MUTATIONS (SID: 22459819) ──
        elif ofd == "ShopeeFood":
            # 1. Item A: OFFLINE_PRICE
            if iid == "3148482929458688":
                row_vals[col_offline_price - 1] = 15000
                row_vals[col_offline_adj - 1] = 3000
                row_vals[col_notes - 1] = "[TEST: OFFLINE_PRICE]"
                print("Applied Shopee Item A: OFFLINE_PRICE Rp 15,000")

            # 2. Item B: PRICE_CHANGE
            elif iid == "3148483367975936":
                row_vals[col_new_fake - 1] = 32000
                row_vals[col_notes - 1] = "[TEST: PRICE_CHANGE NORMAL]"
                print("Applied Shopee Item B: PRICE_CHANGE to Rp 32,000")

            # 3. Item C: NAME_CHANGE
            elif iid == "3148483850695168":
                row_vals[col_item_name_imp - 1] = "Item C Shopee Mantap"
                row_vals[col_notes - 1] = "[TEST: NAME_CHANGE]"
                print("Applied Shopee Item C: NAME_CHANGE")

            # 4. Nasi Goreng Spesial: DESCRIPTION_CHANGE
            elif iid == "3240685897501184":
                row_vals[col_desc - 1] = "Nasi goreng spesial dengan bumbu rahasia unit test."
                row_vals[col_notes - 1] = "[TEST: DESCRIPTION_CHANGE]"
                print("Applied Shopee Nasi Goreng Spesial: DESCRIPTION_CHANGE")

            # 5. Nasi Goreng Cabe Ijo: AVAILABILITY_CHANGE
            elif iid == "3240701593667072":
                row_vals[col_avail - 1] = "Unavailable"
                row_vals[col_notes - 1] = "[TEST: AVAILABILITY_CHANGE]"
                print("Applied Shopee Nasi Goreng Cabe Ijo: AVAILABILITY_CHANGE")

            # 6. Sego Sadukan: CATEGORY_CHANGE
            elif iid == "3240687955772928":
                target_cat = "Category B"
                row_vals[col_cat - 1] = target_cat
                row_vals[col_cat_id - 1] = CATEGORY_ID_MAP[("ShopeeFood", target_cat)]
                row_vals[col_notes - 1] = "[TEST: CATEGORY_CHANGE]"
                print(f"Applied Shopee Sego Sadukan: CATEGORY_CHANGE to '{target_cat}' (ID: {row_vals[col_cat_id - 1]})")

        # ── GOFOOD MUTATIONS (SID: G828092803) ──
        elif ofd == "GoFood":
            # 1. Es Krim Kacang: PHOTO_CHANGE & DESCRIPTION_CHANGE
            if iid == "b032fb90-8d03-45f9-b32b-09bc2ea129e0":
                row_vals[col_photo - 1] = "https://images.unsplash.com/photo-1501443762994-82bd5dace89a"
                row_vals[col_desc - 1] = "Es krim kacang gurih manis racikan unit test."
                row_vals[col_notes - 1] = "[TEST: PHOTO_CHANGE & DESCRIPTION_CHANGE]"
                print("Applied GoFood Es Krim Kacang: PHOTO_CHANGE & DESCRIPTION_CHANGE")

            # 2. Es Krim Kacang Susu: AVAILABILITY_CHANGE
            elif iid == "e7751e20-6887-4050-b737-5518752c8508":
                row_vals[col_avail - 1] = "Unavailable"
                row_vals[col_notes - 1] = "[TEST: AVAILABILITY_CHANGE]"
                print("Applied GoFood Es Krim Kacang Susu: AVAILABILITY_CHANGE")

            # 3. Ayam Goreng Kremes: CATEGORY_CHANGE
            elif iid == "2fc75766-5f76-41e8-94c2-ed6ff491f49a":
                target_cat = "Salad & Sehat"
                row_vals[col_cat - 1] = target_cat
                row_vals[col_cat_id - 1] = CATEGORY_ID_MAP[("GoFood", target_cat)]
                row_vals[col_notes - 1] = "[TEST: CATEGORY_CHANGE]"
                print(f"Applied GoFood Ayam Goreng Kremes: CATEGORY_CHANGE to '{target_cat}' (ID: {row_vals[col_cat_id - 1]})")

            # 4. Ayam Bakar Madu: NAME_CHANGE
            elif iid == "4e40a678-5582-4ca6-8251-8f04380c9f19":
                row_vals[col_item_name_imp - 1] = "Ayam Bakar Madu GoFood Mantap"
                row_vals[col_notes - 1] = "[TEST: NAME_CHANGE]"
                print("Applied GoFood Ayam Bakar Madu: NAME_CHANGE")

            # 5. Ayam Rica-Rica: PRICE_CHANGE
            elif iid == "1e07b8c8-9e64-41b2-a698-5fca6fd6d2c4":
                row_vals[col_new_fake - 1] = 52000
                row_vals[col_notes - 1] = "[TEST: PRICE_CHANGE NORMAL]"
                print("Applied GoFood Ayam Rica-Rica: PRICE_CHANGE to Rp 52,000")

            # 6. Lumpia Udang: PRICE_WARNING_STEP_PUSH
            elif iid == "829a4b6e-4861-4630-a525-5ddeba86ffa2":
                row_vals[col_new_fake - 1] = 58000
                row_vals[col_notes - 1] = "[TEST: PRICE_WARNING_STEP_PUSH]"
                print("Applied GoFood Lumpia Udang: PRICE_WARNING_STEP_PUSH to Rp 58,000")

            # 7. Salad Buah: OFFLINE_PRICE
            elif iid == "6e61bb4f-1c97-45d2-a41d-6ab6075e5e04":
                row_vals[col_offline_price - 1] = 40000
                row_vals[col_offline_adj - 1] = 5000
                row_vals[col_notes - 1] = "[TEST: OFFLINE_PRICE]"
                print("Applied GoFood Salad Buah: OFFLINE_PRICE Rp 40,000")

        rows_data.append(row_vals)

    # ── APPEND NEW ITEMS & NEW CATEGORIES ──
    # 1. Grab New Item & New Category
    new_grab = [None] * max_c
    new_grab[col_ofd - 1] = "GrabFood"
    new_grab[col_outlet - 1] = "Es Degan Jeruk Nipis, Foodnesia"
    new_grab[col_sid - 1] = "6-C7MAC3LDBFCCRE"
    new_grab[col_cat - 1] = "Kategori Spesial Grab"
    new_grab[col_cat_id - 1] = "" # New category has no ID
    new_grab[col_item_id - 1] = "" # New item has no ID
    new_grab[col_item - 1] = "Menu Baru Grab Test"
    new_grab[col_avail - 1] = "Available"
    new_grab[col_new_fake - 1] = 25000
    new_grab[col_curr_fake - 1] = 25000
    new_grab[col_notes - 1] = "[TEST: NEW_ITEM & NEW_CATEGORY]"
    rows_data.append(new_grab)
    print("Added New Item Grab: Menu Baru Grab Test (Kategori Spesial Grab)")

    # 2. Shopee New Item & New Category
    new_shopee = [None] * max_c
    new_shopee[col_ofd - 1] = "ShopeeFood"
    new_shopee[col_outlet - 1] = "Es Degan Jeruk Nipis, Foodnesia"
    new_shopee[col_sid - 1] = "22459819"
    new_shopee[col_cat - 1] = "Kategori Baru Shopee"
    new_shopee[col_cat_id - 1] = "" # New category has no ID
    new_shopee[col_item_id - 1] = "" # New item has no ID
    new_shopee[col_item - 1] = "Ayam Bakar Madu Shopee"
    new_shopee[col_avail - 1] = "Available"
    new_shopee[col_new_fake - 1] = 28000
    new_shopee[col_curr_fake - 1] = 28000
    new_shopee[col_notes - 1] = "[TEST: NEW_ITEM & NEW_CATEGORY]"
    rows_data.append(new_shopee)
    print("Added New Item Shopee: Ayam Bakar Madu Shopee (Kategori Baru Shopee)")

    # 3. GoFood New Item & New Category
    new_gofood = [None] * max_c
    new_gofood[col_ofd - 1] = "GoFood"
    new_gofood[col_outlet - 1] = "WonderFood 3, Sambikerep/Villa Sentra/A6-18"
    new_gofood[col_sid - 1] = "G828092803"
    new_gofood[col_cat - 1] = "Kategori Baru GoFood"
    new_gofood[col_cat_id - 1] = "" # New category has no ID
    new_gofood[col_item_id - 1] = "" # New item has no ID
    new_gofood[col_item - 1] = "Menu Baru GoFood Test"
    new_gofood[col_avail - 1] = "Available"
    new_gofood[col_new_fake - 1] = 35000
    new_gofood[col_curr_fake - 1] = 35000
    new_gofood[col_notes - 1] = "[TEST: NEW_ITEM & NEW_CATEGORY]"
    rows_data.append(new_gofood)
    print("Added New Item GoFood: Menu Baru GoFood Test (Kategori Baru GoFood)")

    # Clear existing data rows
    ws.delete_rows(2, ws.max_row)

    # Insert updated rows
    for r_idx, r_data in enumerate(rows_data, start=2):
        for c_idx, val in enumerate(r_data, start=1):
            if val is not None:
                ws.cell(r_idx, c_idx).value = val

    wb.save(OUTPUT_FILE)
    print(f"Successfully generated: {OUTPUT_FILE} with {len(rows_data)} item rows.")

if __name__ == "__main__":
    generate()

