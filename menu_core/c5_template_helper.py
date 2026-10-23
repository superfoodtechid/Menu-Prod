# -*- coding: utf-8 -*-
"""Helper untuk manajemen template C5 Excel (O. C5 .xlsx).

Menjamin integritas template C5:
1. Memastikan Baris 2 ('Don't touch row 2') tidak pernah disentuh, ditimpa, atau dihapus.
2. Data menu dan modifier selalu dimulai dari Baris 3.
3. Seluruh tab bawaan (Item ME, Modifier ME, Structure, Pricing, Meaning, Tech Hierarchy, dll) tetap utuh.
"""

import os
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple
from openpyxl.styles import Font

DATA_START_ROW = 3
TEMPLATE_RESERVED_ROW = 2

FOLDER_LINK_SHEET_NAME = "Folder Link"
FOLDER_LINK_LABEL_HEADER = "Platform | Listing Name | SID | YYYY-MM-DD HH:MM"
FOLDER_LINK_URL_HEADER = "Folder Link"


def get_c5_template_path(base_dir: Optional[Path] = None) -> Path:
    """Mengembalikan path file template C5 yang valid.
    Memprioritaskan 'O. C5 .xlsx' dan fallback ke 'O. C5 Template.xlsx'.
    """
    if base_dir is None:
        base_dir = Path(__file__).resolve().parents[1]

    primary = base_dir / "O. C5 .xlsx"
    if primary.exists():
        return primary

    fallback = base_dir / "O. C5 Template.xlsx"
    if fallback.exists():
        return fallback

    return primary


def prepare_c5_sheet(sheet) -> None:
    """Membersihkan baris data menu/modifier sebelumnya dari baris 3 ke bawah.
    PENTING: Jangan pernah menghapus baris 2 ('Don't touch row 2') karena menjadi
    acuan rumus ArrayFormula di baris 1.
    """
    if sheet.max_row >= DATA_START_ROW:
        sheet.delete_rows(DATA_START_ROW, sheet.max_row - (DATA_START_ROW - 1))


def get_c5_headers(sheet) -> Dict[str, int]:
    """Mengambil dictionary {nama_header: nomor_kolom} dari baris 1 sheet.
    Hanya mengekstrak nama kolom string murni dan mengabaikan objek ArrayFormula
    agar tidak merusak atau menimpa kolom formula bawaan template.
    """
    headers = {}
    for cell in sheet[1]:
        val = cell.value
        if val is not None and isinstance(val, str):
            clean_val = val.strip()
            if clean_val:
                headers[clean_val] = cell.column
    return headers


def is_c5_reserved_row(row_values) -> bool:
    """Memeriksa apakah suatu baris adalah baris khusus template (Baris 2 'Don't touch row 2')
    sehingga harus dilewati saat membaca data menu.
    """
    if not row_values:
        return True
    first_few = row_values[:5] if isinstance(row_values, (list, tuple)) else [row_values]
    for val in first_few:
        if val is not None and "don't touch" in str(val).lower():
            return True
    return False


def upsert_folder_link_sheet(wb, folder_entries: List[Tuple[str, str]]) -> None:
    """Create or update the 'Folder Link' sheet with folder labels and Google Drive URLs."""
    if FOLDER_LINK_SHEET_NAME in wb.sheetnames:
        ws = wb[FOLDER_LINK_SHEET_NAME]
        if ws.max_row >= 1:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(title=FOLDER_LINK_SHEET_NAME)

    header_font = Font(bold=True)
    c1 = ws.cell(row=1, column=1, value=FOLDER_LINK_LABEL_HEADER)
    c2 = ws.cell(row=1, column=2, value=FOLDER_LINK_URL_HEADER)
    c1.font = header_font
    c2.font = header_font

    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 65

    seen = set()
    row_idx = 2
    for label, url in folder_entries or []:
        clean_label = str(label or "").strip()
        clean_url = str(url or "").strip()
        if not clean_label and not clean_url:
            continue
        dedup_key = (clean_label, clean_url)
        if dedup_key in seen:
            continue
        seen.add(dedup_key)
        ws.cell(row=row_idx, column=1, value=clean_label)
        ws.cell(row=row_idx, column=2, value=clean_url)
        row_idx += 1


def read_folder_link_entries(wb_or_path) -> List[Tuple[str, str]]:
    """Read (label, url) rows from the 'Folder Link' sheet of an O.C5 workbook or path."""
    import openpyxl

    close_after = False
    wb = wb_or_path
    if isinstance(wb_or_path, (str, Path)):
        if not os.path.exists(str(wb_or_path)):
            return []
        try:
            wb = openpyxl.load_workbook(str(wb_or_path), data_only=True, read_only=True)
            close_after = True
        except Exception:
            return []

    entries: List[Tuple[str, str]] = []
    try:
        if FOLDER_LINK_SHEET_NAME in wb.sheetnames:
            ws = wb[FOLDER_LINK_SHEET_NAME]
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if r_idx == 1 or not row:
                    continue
                label = str(row[0] if len(row) > 0 and row[0] is not None else "").strip()
                url = str(row[1] if len(row) > 1 and row[1] is not None else "").strip()
                if label or url:
                    entries.append((label, url))
    except Exception:
        pass
    finally:
        if close_after and wb is not None:
            try:
                wb.close()
            except Exception:
                pass
    return entries


