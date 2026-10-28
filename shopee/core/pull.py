# -*- coding: utf-8 -*-
import os
import sys
import json
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
import pandas as pd
from pathlib import Path
from shopee.core.client import ShopeeClient, ShopeeModifyClient

# Dynamic paths based on current file location
FILE_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = FILE_DIR.parent.parent
MENU_DIR = WORKSPACE_DIR / "shopee"
AUTOMATION_DIR = WORKSPACE_DIR / "src" / "shopee-omzet-automation"
if str(AUTOMATION_DIR) not in sys.path:
    sys.path.insert(0, str(AUTOMATION_DIR))
from core import browser
from shopee.core.session_utils import (
    get_fast_shopee_session,
    invalidate_shopee_session,
    resolve_shopee_merchant_id,
    resolve_shopee_outlet_context,
)
from menu_core.c5_template_helper import (
    get_c5_template_path,
    prepare_c5_sheet,
    get_c5_headers,
    DATA_START_ROW,
    is_c5_reserved_row,
    upsert_folder_link_sheet,
    FOLDER_LINK_SHEET_NAME,
    FOLDER_LINK_LABEL_HEADER,
    FOLDER_LINK_URL_HEADER,
)
from shopee.core.push_c5 import is_valid_shopee_image_hash

IMG_BASE = "https://down-id.img.susercontent.com/file"

def list_menu_shopee(store_metadata: dict) -> tuple[bool, list | str]:
    from shopee.core.item.edit import _boot_client
    client, err = _boot_client(store_metadata, headless=None)
    if not client:
        return False, f"Boot client failed: {err}"
        
    store_id = store_metadata["store_id"]
    catalogs = client.get_store_dishes(store_id)
    return True, catalogs

def get_shopee_master_credentials() -> tuple[str, str]:
    """
    Membaca kredensial master Shopee dari environment variable (.env)
    atau dari file credentials.json di root project.
    """
    user = os.getenv("SHOPEE_MASTER_USERNAME") or os.getenv("SHOPEE_USERNAME")
    pwd = os.getenv("SHOPEE_MASTER_PASSWORD") or os.getenv("SHOPEE_PASSWORD")
    if user and pwd:
        return user.strip(), pwd.strip()

    cred_file = WORKSPACE_DIR / "credentials.json"
    if cred_file.exists():
        try:
            with open(cred_file, "r", encoding="utf-8") as f:
                data = json.load(f)
                s_cred = data.get("shopee", {})
                u = s_cred.get("master_username") or s_cred.get("username")
                p = s_cred.get("master_password") or s_cred.get("password")
                if u and p:
                    return str(u).strip(), str(p).strip()
        except Exception:
            pass

    return "allvbadmin", "Master!00!"


def _attempt_dishes_pull_recovery(
    client: ShopeeClient,
    store_id: str,
    merchant_id: str,
    target_name: str,
    extra_cookies: dict
) -> tuple[list[dict], str, str]:
    """
    Pemulihan otomatis saat API dishes merespons 1130014 (authority invalid)
    atau 1130002 (need to select a store).
    """
    catalogs = []
    current_sid = store_id
    current_mid = merchant_id
    orig_sid_cookie = extra_cookies.get("shopee_tob_entity_id")
    orig_mid_cookie = extra_cookies.get("shopee_foody_mid")
    orig_client_eid = client.entity_id
    orig_client_mid = client.merchant_id

    # 1. Coba resolusi ulang store ID dan merchant ID dari database/cache
    rec_sid, rec_mid = resolve_shopee_outlet_context(
        {"store_id": current_sid, "merchant_id": current_mid, "nama_portal": target_name}
    )
    if not rec_sid or rec_sid == current_sid:
        rec_sid, rec_mid = resolve_shopee_outlet_context(
            {"nama_portal": target_name, "merchant_name": target_name}
        )

    if rec_sid and (rec_sid != current_sid or rec_mid != current_mid):
        print(f"[RECOVERY] Mengoreksi store context: sid={current_sid}->{rec_sid}, mid={current_mid}->{rec_mid} (error {client.last_code})")
        current_sid = rec_sid
        if rec_mid:
            current_mid = rec_mid
            extra_cookies["shopee_foody_mid"] = current_mid
            client.merchant_id = current_mid
        extra_cookies["shopee_tob_entity_id"] = str(current_sid)
        client.entity_id = current_sid
        catalogs = client.get_store_dishes(current_sid)

    # 2. Jika masih gagal, query daftar store resmi yang diotorisasi akun via API
    if not catalogs and client.last_code in (1130014, 1130002):
        try:
            stores = client.get_stores()
            if stores:
                matched_sid = None
                for s in stores:
                    s_name = str(s.get("name", "")).lower()
                    s_id = str(s.get("id", "")).strip()
                    if target_name.lower() in s_name or s_name in target_name.lower():
                        matched_sid = s_id
                        break

                if matched_sid and matched_sid != current_sid:
                    print(f"[RECOVERY] Menggunakan Store ID resmi dari get_stores: {current_sid} -> {matched_sid}")
                    current_sid = matched_sid
                    extra_cookies["shopee_tob_entity_id"] = str(current_sid)
                    client.entity_id = current_sid
                    catalogs = client.get_store_dishes(current_sid)
        except Exception as store_err:
            print(f"[-] Warning: Gagal query get_stores untuk recovery: {store_err}")

    if catalogs:
        return catalogs, current_sid, current_mid

    # Jika pemulihan gagal menemukan katalog, kembalikan state awal tanpa mengotori cookie / client
    if orig_sid_cookie is not None:
        extra_cookies["shopee_tob_entity_id"] = orig_sid_cookie
    else:
        extra_cookies.pop("shopee_tob_entity_id", None)

    if orig_mid_cookie is not None:
        extra_cookies["shopee_foody_mid"] = orig_mid_cookie
    else:
        extra_cookies.pop("shopee_foody_mid", None)

    client.entity_id = orig_client_eid
    client.merchant_id = orig_client_mid
    return [], store_id, merchant_id


def extract_shopee_menu(store_metadata: dict, output_dir: str, headless: bool = True):
    # Resolusi Store ID fisik dan Group ID (Merchant ID) secara presisi di awal
    resolved_sid, resolved_mid = resolve_shopee_outlet_context(store_metadata)
    if resolved_sid:
        store_metadata['store_id'] = resolved_sid
    if resolved_mid:
        store_metadata['merchant_id'] = resolved_mid

    store_id = store_metadata.get('store_id', '')
    if isinstance(store_id, str):
        store_id = store_id.strip().split('.')[0]
    if not store_id or store_id == '-' or store_id.lower() == 'nan':
        store_id = None
        
    nama_portal = store_metadata.get('nama_portal')
    if nama_portal and str(nama_portal).strip() and str(nama_portal).lower() not in ('nan', '-', 'none'):
        target_name = str(nama_portal).strip()
    else:
        m_name = store_metadata.get('merchant_name', '')
        if not m_name or m_name.lower() == 'nan' or m_name == '-':
            target_name = store_metadata.get('nama_resto_final') or store_metadata.get('nama_outlet') or ''
        else:
            target_name = m_name
        
    nama_panjang = store_metadata.get('nama_resto_final') or store_metadata.get('nama_outlet') or target_name
    nama_pendek = store_metadata.get('brand') or store_metadata.get('nama_pendek') or store_metadata.get('nama_outlet') or target_name
    
    master_user, master_pass = get_shopee_master_credentials()
    username = master_user
    password = master_pass
    # Penarikan menu (PULL) selalu menggunakan akun master allvbadmin
    session_file = AUTOMATION_DIR / "data" / f"session_{master_user}.json"
    if not session_file.exists():
        session_file = MENU_DIR / "data" / f"session_{master_user}.json"
        if not session_file.exists():
            session_file = WORKSPACE_DIR / "data" / f"session_{master_user}.json"
    browser.set_session_file(session_file)
    active_account = master_user
            
    job_id = store_metadata.get('job_id')
    def _is_cancelled():
        if not job_id:
            return False
        try:
            from menu_core.job_control import is_job_cancelled
            return is_job_cancelled(job_id)
        except Exception:
            return False

    if _is_cancelled():
        print(f"[!] Shopee job {job_id} telah dibatalkan oleh pengguna.")
        return False, "user membatalkan otp"

    session_data = None
    if store_id:
        fast_meta = dict(store_metadata)
        fast_meta.setdefault("username", master_user)
        fast_meta.setdefault("store_id", store_id)
        fast_meta.setdefault("nama_portal", target_name)
        fast_sess = get_fast_shopee_session(
            fast_meta, base_dir=WORKSPACE_DIR, require_entity_match=True
        )
        if fast_sess and fast_sess.get("shopee_tob_token"):
            print(f"⚡ [FAST-PATH PULL] Menggunakan sesi master Shopee aktif untuk '{master_user}' (~0.2s).")
            session_data = fast_sess

    if not session_data:
        print(f"[*] Membuka browser (headless={headless}) untuk akun '{active_account}' (target: '{username}'), merchant: '{target_name}', session: '{session_file.name}'...")
        try:
            session_data = browser.get_session(
                username=active_account,
                password=password,
                headless=headless,
                close_browser=True,
                target_name=target_name,
                target_store_id=store_id,
                interactive=True,
                strict_entity_context=True,
            )
        except Exception as e:
            if "user membatalkan otp" in str(e).lower() or _is_cancelled():
                return False, "user membatalkan otp"
            return False, f"Gagal menginisialisasi browser: {e}"

    if _is_cancelled():
        print(f"[!] Shopee job {job_id} telah dibatalkan setelah inisialisasi browser.")
        return False, "user membatalkan otp"
    
    if not session_data or "shopee_tob_token" not in session_data:
        return False, "Gagal menginisialisasi browser atau memilih merchant."

    session_entity_id = str(session_data.get("shopee_tob_entity_id") or "").strip()
    if store_id and session_entity_id != str(store_id).strip():
        return False, (
            f"Sesi Shopee tidak cocok dengan outlet target (store_id={store_id}, "
            f"session_entity_id={session_entity_id or 'kosong'}); token tidak digunakan."
        )
        
    tob_token = session_data["shopee_tob_token"]
    extra_cookies = session_data.get("extra_cookies", {})
    
    # If store_id was not originally known, resolve it from the active session
    if not store_id:
        session_store_id = session_data.get("shopee_tob_entity_id")
        if session_store_id:
            store_id = str(session_store_id).strip()
            store_metadata['store_id'] = store_id
            print(f"[+] Berhasil mendeteksi Store ID secara dinamis: {store_id}")
            
            try:
                db_path = MENU_DIR / "data" / "menu_management.db"
                import sqlite3
                conn = sqlite3.connect(str(db_path))
                cursor = conn.cursor()
                cursor.execute(
                    "INSERT INTO outlets (store_id, platform, username, password, merchant_name) "
                    "VALUES (?, 'shopee', ?, ?, ?) "
                    "ON CONFLICT(store_id) DO UPDATE SET username=excluded.username, password=excluded.password, merchant_name=excluded.merchant_name",
                    (store_id, username, password, target_name)
                )
                conn.commit()
                conn.close()
                print(f"[+] Menyimpan Store ID {store_id} untuk '{target_name}' ke database lokal.")
            except Exception as dbe:
                print(f"[-] Warning: Gagal menyimpan Store ID ke DB: {dbe}")
        else:
            return False, "Gagal mendeteksi Store ID untuk merchant ini."
            
    real_merchant_id = resolve_shopee_merchant_id(
        store_metadata,
        session_data=session_data,
        tob_token=tob_token
    )
    if not real_merchant_id and session_data.get("merchant_id"):
        real_merchant_id = str(session_data["merchant_id"]).strip()

    extra_cookies = session_data.get("extra_cookies", {}).copy()
    extra_cookies["shopee_user_name"] = username
    extra_cookies["shopee_request_from"] = "partner_web"
    if real_merchant_id:
        extra_cookies["shopee_foody_mid"] = real_merchant_id
    if store_id:
        extra_cookies["shopee_tob_entity_id"] = str(store_id)

    session_source = session_data.get("source", "browser")
    print(f"[PULL SESSION] source={session_source}, store_id={store_id}, merchant_id={real_merchant_id}")
    
    client = ShopeeClient(
        tob_token=tob_token,
        entity_id=store_id,
        extra_cookies=extra_cookies,
        merchant_id=real_merchant_id,
        username=username
    )
        
    try:
        print(f"[*] Menarik data menu ShopeeFood untuk: {target_name} ({store_id})...")
        # The catalog API is authoritative for PULL; a separate GetUserInfo
        # preflight can reject valid partner sessions and has different context.
        catalogs = client.get_store_dishes(store_id)
        if not catalogs and client.last_code in (1130014, 1130002):
            rec_cats, rec_sid, rec_mid = _attempt_dishes_pull_recovery(
                client=client,
                store_id=store_id,
                merchant_id=real_merchant_id,
                target_name=target_name,
                extra_cookies=extra_cookies
            )
            if rec_cats:
                catalogs = rec_cats
                store_id = rec_sid
                real_merchant_id = rec_mid

        if not catalogs:
            is_auth_error = client.last_http_status in (401, 403) or client.last_code in (
                3004, 1000002, 1000003, 1130001
            )
            if is_auth_error:
                print(f"[WARN] Sesi Shopee ditolak (code: {client.last_code}, reason: '{client.last_error}') untuk '{username}', store_id={store_id}. Membersihkan cache dan mencoba satu pemulihan browser...")
                try:
                    if session_file.exists():
                        session_file.unlink()
                        print(f"[+] Session file cache '{session_file.name}' berhasil dihapus.")
                except Exception as del_err:
                    print(f"[-] Gagal menghapus session file: {del_err}")

                invalidate_shopee_session({"username": master_user, "store_id": store_id})
            should_refresh_browser = is_auth_error
            if should_refresh_browser:
                session_data = None
                print(f"[PULL SESSION] API katalog menolak sesi secara eksplisit (code={client.last_code}, HTTP={client.last_http_status}); mencoba satu pemulihan browser untuk '{target_name}'.")
                session_data = browser.get_session(
                    username=username,
                    password=password,
                    headless=headless,
                    close_browser=True,
                    target_name=target_name,
                    target_store_id=store_id,
                    interactive=True,
                    strict_entity_context=True,
                    force_token_refresh=True,
                )
            else:
                session_data = None
                print(f"[PULL SESSION] Katalog gagal tanpa penolakan autentikasi eksplisit (code={client.last_code}, HTTP={client.last_http_status}); tidak melakukan relogin otomatis.")

            if session_data and store_id and str(session_data.get("shopee_tob_entity_id") or "").strip() != str(store_id).strip():
                print(f"[PULL SESSION] Pemulihan menghasilkan entity yang tidak cocok; retry katalog dibatalkan untuk store_id={store_id}.")
                session_data = None
                client.last_code = 1130001
                client.last_error = "token/session entity mismatch after recovery"

            if session_data and "shopee_tob_token" in session_data:
                tob_token = session_data["shopee_tob_token"]
                resolved_sid, resolved_mid = resolve_shopee_outlet_context(
                    store_metadata,
                    session_data=session_data
                )
                new_mid = resolved_mid or real_merchant_id
                extra_cookies = session_data.get("extra_cookies", {}).copy()
                extra_cookies["shopee_user_name"] = username
                extra_cookies["shopee_request_from"] = "partner_web"
                if new_mid:
                    extra_cookies["shopee_foody_mid"] = new_mid
                if store_id:
                    extra_cookies["shopee_tob_entity_id"] = str(store_id)

                client = ShopeeClient(
                    tob_token=tob_token,
                    entity_id=store_id,
                    extra_cookies=extra_cookies,
                    merchant_id=new_mid,
                    username=username
                )
                print(f"[*] Mencoba ulang endpoint katalog satu kali untuk store_id={store_id}...")
                catalogs = client.get_store_dishes(store_id)
                if not catalogs and client.last_code in (1130014, 1130002):
                    rec_cats, rec_sid, rec_mid = _attempt_dishes_pull_recovery(
                        client=client,
                        store_id=store_id,
                        merchant_id=new_mid,
                        target_name=target_name,
                        extra_cookies=extra_cookies
                    )
                    if rec_cats:
                        catalogs = rec_cats
                        store_id = rec_sid
                        new_mid = rec_mid

        if not catalogs:
            err_details = f"code={client.last_code}, msg={client.last_error}" if client.last_code else "katalog kosong"
            return False, f"Tidak ada data catalog/dishes yang ditemukan ({err_details})."
            
        print(f"[*] Ditemukan {len(catalogs)} kategori menu.")

        # Save raw Shopee catalog snapshot to shopee/API/menu-response-<store_id>.json for baseline comparison
        try:
            shopee_api_dir = WORKSPACE_DIR / "shopee" / "API"
            shopee_api_dir.mkdir(parents=True, exist_ok=True)
            snapshot_path = shopee_api_dir / f"menu-response-{store_id}.json"
            snapshot_path.write_text(json.dumps({"data": {"catalogs": catalogs}}, indent=2), encoding="utf-8")
            # Also keep default snapshot fallback updated
            (shopee_api_dir / "menu-response.json").write_text(json.dumps({"data": {"catalogs": catalogs}}, indent=2), encoding="utf-8")
            print(f"   💾 Snapshot menu Shopee berhasil disimpan ke: {snapshot_path}")
        except Exception as e:
            print(f"   ⚠️ Gagal menyimpan snapshot menu Shopee: {e}")

        all_dishes = []
        dish_ids_with_modifiers = []
        
        for cat in catalogs:
            cat_name = cat.get('name', 'Menu Lainnya')
            cat_id = str(cat.get('id', ''))
            dishes = cat.get('dishes', [])
            for dish in dishes:
                dish_id = str(dish.get('id'))
                dish_name = dish.get('name', '')
                price_raw = dish.get('price', '0')
                list_price_raw = dish.get('list_price', '0')
                description = dish.get('description', '')
                available = dish.get('available', True)
                opt_group_count = dish.get('option_group_count', 0)
                sales_volume = dish.get('sales_volume', 0)
                picture = dish.get('picture', '')
                discount_pct = float(dish.get('discount_percentage', 0))
                
                price = float(price_raw) / 100000.0
                list_price = float(list_price_raw) / 100000.0 if (list_price_raw and float(list_price_raw) > 0) else price
                
                if list_price > price:
                    current_fake_price = list_price
                    current_real_price = price
                    pct_val = round(((list_price - price) / list_price) * 100.0, 2)
                    current_slash_pct = f"{int(pct_val)}%" if pct_val.is_integer() else f"{pct_val}%"
                    current_slash_rp = list_price - price
                else:
                    current_fake_price = list_price
                    current_real_price = list_price
                    current_slash_pct = "0%"
                    current_slash_rp = 0
                
                available_str = "Available" if available else "Unavailable"
                if picture and (str(picture).startswith(("http://", "https://")) or is_valid_shopee_image_hash(picture)):
                    picture_url = str(picture) if str(picture).startswith(("http://", "https://")) else f"{IMG_BASE}/{picture}"
                else:
                    picture_url = ""
                link_outlet = f"https://shopee.co.id/now-food/shop/{store_id}"
                
                dish_info = {
                    'ofd': 'ShopeeFood',
                    'link_outlet': link_outlet,
                    'nama_panjang': nama_panjang,
                    'nama_pendek': nama_pendek,
                    'store_id': store_id,
                    'cat_id': cat_id,
                    'nama_kategori': cat_name,
                    'nama_item': dish_name,
                    'jumlah_terjual': sales_volume,
                    'opt_group_count': opt_group_count,
                    'deskripsi_item': description,
                    'keyword': '',
                    'harga_sebelum_promo': current_fake_price,
                    'harga_setelah_promo': current_real_price,
                    'discount_pct_val': current_slash_pct,
                    'slash_price_rp': current_slash_rp,
                    'ketersediaan': available_str,
                    'visibility': 'Show',
                    'link_foto': picture_url,
                    'dish_id': dish_id,
                    'jumlah_modifier_group': 0,
                    'jumlah_modifier': 0
                }
                all_dishes.append(dish_info)
                
                if opt_group_count > 0:
                    dish_ids_with_modifiers.append(dish_id)

        print(f"[*] Total {len(all_dishes)} item ditemukan.")
        print(f"[*] Menarik modifier untuk {len(dish_ids_with_modifiers)} item yang memiliki topping/opsi...")
        
        dish_to_opt_groups = {}
        if dish_ids_with_modifiers:
            def _fetch_group(did):
                try:
                    return did, client.get_store_option_groups(store_id, dish_ids=[did])
                except Exception as e:
                    print(f"[!] Error fetching option groups for dish {did}: {e}")
                    return did, []

            max_workers = min(10, len(dish_ids_with_modifiers))
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = {executor.submit(_fetch_group, did): did for did in dish_ids_with_modifiers}
                for f in as_completed(futures):
                    did, res = f.result()
                    dish_to_opt_groups[did] = res

        modifier_rows = []
        for dish_id in dish_ids_with_modifiers:
            dish_obj = next((d for d in all_dishes if d['dish_id'] == dish_id), None)
            if not dish_obj:
                continue
                
            opt_groups = dish_to_opt_groups.get(dish_id, [])
            dish_obj['jumlah_modifier_group'] = len(opt_groups)
            total_modifiers_count = 0
            
            group_names = []
            for group in opt_groups:
                opt_group_info = group.get('option_group', {})
                group_name = opt_group_info.get('name', '').strip()
                if group_name:
                    group_names.append(group_name)
                select_min = opt_group_info.get('select_min', 0)
                select_max = opt_group_info.get('select_max', 0)
                options = group.get('options', [])
                
                total_modifiers_count += len(options)
                tipe_modifier = "Pilihan Tunggal" if select_max == 1 else "Pilihan Ganda"
                
                for opt in options:
                    opt_name = opt.get('name', '')
                    opt_price = float(opt.get('price', '0')) / 100000.0
                    opt_available = opt.get('available', True)
                    opt_available_str = "Available" if opt_available else "Unavailable"
                    
                    modifier_rows.append({
                        'ofd': 'ShopeeFood',
                        'outlet_name': nama_panjang,
                        'outlet_short_name': nama_pendek,
                        'outlet_link': dish_obj['link_outlet'],
                        'sid': store_id,
                        'item': dish_obj['nama_item'],
                        'modifier_group_id': str(opt_group_info.get('id', '')),
                        'modifier_group': group_name,
                        'modifier_id': str(opt.get('id', '')),
                        'modifier': opt_name,
                        'min': select_min,
                        'max': select_max,
                        'availability': opt_available_str,
                        'visibility': 'Show',
                        'current_price': opt_price
                    })
            
            dish_obj['jumlah_modifier'] = total_modifiers_count
            dish_obj['keyword'] = ", ".join(group_names)
            
        item_cols = [
            'OFD', 'Outlet Name', 'Outlet Short Name', 'Outlet Link', 'SID',
            'Category ID', 'Category', 'Item ID', 'Item', 'Photo Link',
            'Description', 'Keyword', 'Total Sold', 'Total Modifier Group', 'Total Modifier',
            'Availability', 'Visibility', 'Current Fake Price (Rp)', 'Current Real Price (Rp)',
            'Current Slash Price (%)', 'Current Slash Price (Rp)', 'New Markup (%)',
            'New Real Price (Rp)', 'Adjustment (Rp)', 'New Final Real Price (Rp)',
            'New Slash Price (%)', 'New Fake Price (Rp)', 'Notes'
        ]
        
        item_data = []
        for d in all_dishes:
            item_data.append([
                d['ofd'], d['nama_panjang'], d['nama_pendek'], d['link_outlet'], d['store_id'],
                d['cat_id'], d['nama_kategori'], d['dish_id'], d['nama_item'], d['link_foto'],
                d['deskripsi_item'], d['keyword'], d['jumlah_terjual'], d['jumlah_modifier_group'], d['jumlah_modifier'],
                d['ketersediaan'], d['visibility'], d['harga_sebelum_promo'], d['harga_setelah_promo'],
                d['discount_pct_val'], d['slash_price_rp'],
                "", "", "", "", "", "", ""
            ])
            
        df_items = pd.DataFrame(item_data, columns=item_cols)
        
        mod_cols = [
            'OFD', 'Outlet Name', 'Outlet Short Name', 'Outlet Link', 'SID',
            'Item', 'Modifier Group ID', 'Modifier Group', 'Modifier ID',
            'Modifier', 'Min', 'Max', 'Availability', 'Visibility', 'Current Price (Rp)'
        ]
        
        mod_data = []
        for m in modifier_rows:
            mod_data.append([
                m['ofd'], m['outlet_name'], m['outlet_short_name'], m['outlet_link'], m['sid'],
                m['item'], m['modifier_group_id'], m['modifier_group'], m['modifier_id'],
                m['modifier'], m['min'], m['max'], m['availability'], m['visibility'], m['current_price']
            ])
            
        df_mods = pd.DataFrame(mod_data, columns=mod_cols)
        os.makedirs(output_dir, exist_ok=True)
        
        import re
        def clean_name(s):
            cleaned = "".join(c for c in s if c.isalnum() or c in (' ', '_', '-')).rstrip()
            return cleaned.replace(' ', '_')
            
        safe_merchant = clean_name(target_name)
        branch_raw = store_metadata.get('brand') or store_metadata.get('nama_resto_final') or store_metadata.get('nama_outlet') or ""
        safe_branch = clean_name(branch_raw)
        
        if safe_branch.lower() == safe_merchant.lower() or not safe_branch:
            combined_name = safe_merchant
        else:
            combined_name = f"{safe_merchant}_{safe_branch}"
            
        combined_name = re.sub(r'_+', '_', combined_name)
        
        # Build paths conforming to "O.C5 {nama_outlet} - {brand}.xlsx"
        raw_outlet = store_metadata.get('nama_outlet') or store_metadata.get('nama_resto_final') or target_name or 'unknown'
        raw_brand = store_metadata.get('brand') or ''
        
        def clean_filename_part(s):
            return "".join(c for c in s if c.isalnum() or c in (' ', '_', '-')).strip()
            
        raw_resto_final = store_metadata.get('nama_resto_final') or ''
        clean_outlet = clean_filename_part(raw_outlet)
        clean_brand = clean_filename_part(raw_brand)
        clean_resto = clean_filename_part(raw_resto_final)
        
        if clean_resto and clean_resto.lower() not in (clean_outlet.lower(), clean_brand.lower()):
            excel_filename = f"O.C5 {clean_resto} - {clean_brand}.xlsx"
        elif clean_brand and clean_brand.lower() != clean_outlet.lower():
            excel_filename = f"O.C5 {clean_outlet} - {clean_brand}.xlsx"
        else:
            excel_filename = f"O.C5 {clean_outlet}.xlsx"
            
        excel_path = os.path.join(output_dir, excel_filename)
        
        # Write to O. C5 Template Excel
        template_path = str(get_c5_template_path(WORKSPACE_DIR))
        import openpyxl

        # ── Kolom yang bisa diisi user secara manual (akan di-preserve saat re-pull) ──
        USER_EDITABLE_COLS = [
            'Keyword',                   # Kol L
            'Item Name Improvement',     # Kol M
            'Design Improvement',        # Kol N
            'Offline Price (Rp)',        # Kol W
            'Offline Adjustment (Rp)',   # Kol X
            'Current Adjustment (Rp)',   # Kol AC
            'New Markup (%)',            # Kol AI
            'New Adjustment (Rp)',       # Kol AK
            'New Slash Price (%)',       # Kol AM
            'Notes',                     # Kol AP
        ]

        # ── Load data user dari C5 lama (jika ada) ──
        prev_user_data = {}  # { item_id: { col_name: value } }
        if os.path.exists(excel_path):
            try:
                wb_prev = openpyxl.load_workbook(excel_path, data_only=True)
                if 'Item' in wb_prev.sheetnames:
                    sheet_prev = wb_prev['Item']
                    prev_headers = {}
                    for cell in sheet_prev[1]:
                        if cell.value:
                            prev_headers[str(cell.value).strip()] = cell.column - 1  # 0-indexed
                    item_id_idx = prev_headers.get('Item ID')
                    edit_col_indices = {col: prev_headers[col] for col in USER_EDITABLE_COLS if col in prev_headers}
                    if item_id_idx is not None:
                        for row_cells in sheet_prev.iter_rows(min_row=DATA_START_ROW, values_only=True):
                            if is_c5_reserved_row(row_cells):
                                continue
                            raw_iid = row_cells[item_id_idx] if item_id_idx < len(row_cells) else None
                            if raw_iid is None or str(raw_iid).strip() in ('', 'None', 'nan'):
                                continue
                            iid = str(raw_iid).strip()
                            row_data = {}
                            for col_name, col_idx in edit_col_indices.items():
                                if col_idx < len(row_cells):
                                    val = row_cells[col_idx]
                                    if val is not None and str(val).strip() not in ('', 'None', 'nan'):
                                        row_data[col_name] = val
                            if row_data:
                                prev_user_data[iid] = row_data
                wb_prev.close()
                if prev_user_data:
                    print(f"[🔄] Ditemukan C5 sebelumnya - memuat {len(prev_user_data)} baris data user untuk di-preserve.")
            except Exception as prev_err:
                print(f"[⚠️] Gagal membaca C5 sebelumnya untuk merge: {prev_err}")

        from menu_core.gdrive_service import format_photo_folder_label, prepare_photo_folder
        owner_name = store_metadata.get("owner") or "Default Owner"
        listing_name = raw_resto_final or raw_outlet or "ShopeeFood Outlet"
        sid = str(store_metadata.get("store_id") or "").strip()
        folder_label = format_photo_folder_label("ShopeeFood", listing_name, sid)
        photo_sync_meta = None

        try:
            wb = openpyxl.load_workbook(template_path)
            
            # 1. Fill Item Sheet
            sheet_item = wb['Item']
            prepare_c5_sheet(sheet_item)
                
            headers_item = get_c5_headers(sheet_item)
            
            new_row_idx = DATA_START_ROW
            for d in all_dishes:
                # Lookup data user lama berdasarkan Item ID (dish_id)
                dish_id_str = str(d.get('dish_id', '')).strip()
                prev_row = prev_user_data.get(dish_id_str, {})

                photo_link_val = d['link_foto']

                mapping = {
                    'OFD': d['ofd'],
                    'Listing Name': d['nama_panjang'],
                    'Outlet Name': d['nama_panjang'],
                    'Outlet Short Name': d['nama_pendek'],
                    'Outlet Link': d['link_outlet'],
                    'SID': d['store_id'],
                    'Category ID': d['cat_id'],
                    'Category': d['nama_kategori'],
                    'Item ID': d['dish_id'],
                    'Item': d['nama_item'],
                    'Photo Link': photo_link_val,
                    'Description': d['deskripsi_item'],
                    'Keyword': d['keyword'],
                    'Total Sold': d['jumlah_terjual'],
                    'Total Modifier Group': d['jumlah_modifier_group'],
                    'Total Modifier': d['jumlah_modifier'],
                    'Availability': d['ketersediaan'],
                    'Visibility': d['visibility'],
                    'Current Fake Price (Rp)': d['harga_sebelum_promo'],
                    'Current Real Price (Rp)': d['harga_setelah_promo'],
                    'Current Slash Price (%)': d['discount_pct_val'],
                    'Current Slash Price (Rp)': d['slash_price_rp']
                }
                
                # Tambahkan kolom user-editable dari C5 lama jika ada
                for col_name in USER_EDITABLE_COLS:
                    if col_name not in mapping and col_name in prev_row:
                        mapping[col_name] = prev_row[col_name]
                
                for key, val in mapping.items():
                    if key in headers_item:
                        col_idx = headers_item[key]
                        if pd.isna(val):
                            val = ""
                        elif key in ['SID', 'Category ID', 'Item ID']:
                            if isinstance(val, float):
                                val = str(int(val)) if val.is_integer() else str(val)
                            else:
                                val = str(val)
                        # Jika kolom ini bisa diisi user & nilai baru kosong, restore nilai lama
                        if key in USER_EDITABLE_COLS and (val == "" or val is None):
                            val = prev_row.get(key, val)
                        sheet_item.cell(row=new_row_idx, column=col_idx, value=val)
                new_row_idx += 1
                        
            # 2. Fill Modifier Sheet
            sheet_mod = wb['Modifier']
            prepare_c5_sheet(sheet_mod)
                
            headers_mod = get_c5_headers(sheet_mod)
            
            new_row_idx = DATA_START_ROW
            for m in modifier_rows:
                mapping_mod = {
                    'OFD': m['ofd'],
                    'Listing Name': m['outlet_name'],
                    'Outlet Name': m['outlet_name'],
                    'Outlet Short Name': m['outlet_short_name'],
                    'Outlet Link': m['outlet_link'],
                    'SID': m['sid'],
                    'Item': m['item'],
                    'Modifier Group ID': m['modifier_group_id'],
                    'Modifier Group': m['modifier_group'],
                    'Modifier ID': m['modifier_id'],
                    'Modifier': m['modifier'],
                    'Min': m['min'],
                    'Max': m['max'],
                    'Availability': m['availability'],
                    'Visibility': m['visibility'],
                    'Current Price (Rp)': m['current_price']
                }
                
                for key, val in mapping_mod.items():
                    if key in headers_mod:
                        col_idx = headers_mod[key]
                        if pd.isna(val):
                            val = ""
                        elif key in ['SID', 'Modifier Group ID', 'Modifier ID', 'Item']:
                            if isinstance(val, float):
                                val = str(int(val)) if val.is_integer() else str(val)
                            else:
                                val = str(val)
                        sheet_mod.cell(row=new_row_idx, column=col_idx, value=val)
                new_row_idx += 1

            photo_sync_meta = None
            try:
                folder_url, photo_sync_meta = prepare_photo_folder(
                    owner_name=owner_name,
                    platform="ShopeeFood",
                    listing_name=listing_name,
                    sid=sid,
                    dishes=[{"item_name": d.get("nama_item"), "photo_url": d.get("link_foto"), "item_id": d.get("dish_id")} for d in all_dishes],
                    excel_path=excel_path,
                    job_id=store_metadata.get("job_id"),
                    workbook=wb,
                    defer_photo_upload=True,
                )
            except Exception as e_drive:
                print(f"[⚠️] Gagal menyiapkan folder foto ShopeeFood: {e_drive}")
                folder_entries = [(folder_label, "")]
                upsert_folder_link_sheet(wb, folder_entries)

            wb.save(excel_path)
            wb.close()
            if prev_user_data:
                print(f"[+] Berhasil menyimpan file catalog dengan data user yang di-preserve ke: {excel_path}")
            else:
                print(f"[+] Berhasil menyimpan file catalog menggunakan template O.C5 ke: {excel_path}")

        except Exception as ex_err:
            print(f"[-] Gagal menulis ke template O.C5: {ex_err}. Fallback ke Excel biasa.")
            # Fallback to standard excel writer if template fails
            if photo_sync_meta is None:
                try:
                    _, photo_sync_meta = prepare_photo_folder(
                        owner_name=owner_name,
                        platform="ShopeeFood",
                        listing_name=listing_name,
                        sid=sid,
                        dishes=[{"item_name": d.get("nama_item"), "photo_url": d.get("link_foto"), "item_id": d.get("dish_id")} for d in all_dishes],
                        excel_path=excel_path,
                        job_id=store_metadata.get("job_id"),
                        skip_excel_update=True,
                        defer_photo_upload=True,
                    )
                except Exception as e_drive:
                    print(f"[⚠️] Gagal menyiapkan folder foto ShopeeFood fallback: {e_drive}")
            resolved_folder_url = (photo_sync_meta or {}).get("drive_folder_url") or ""
            df_folder_links = pd.DataFrame(
                [[folder_label, resolved_folder_url]],
                columns=[FOLDER_LINK_LABEL_HEADER, FOLDER_LINK_URL_HEADER],
            )
            with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
                df_items.to_excel(writer, sheet_name='Item', index=False)
                df_mods.to_excel(writer, sheet_name='Modifier', index=False)
                df_folder_links.to_excel(writer, sheet_name=FOLDER_LINK_SHEET_NAME, index=False)

        return True, {
            'items_csv': None,
            'mods_csv': None,
            'excel': excel_path,
            'items_count': len(df_items),
            'mods_count': len(df_mods),
            'photo_sync': photo_sync_meta,
        }
    except Exception as e:
        return False, f"Error selama ekstraksi menu: {e}"


def run_shopee_extraction_worker(store_metadata: dict, output_dir: str, result_queue) -> None:
    """Run Shopee extraction in a killable worker process."""
    try:
        headless = bool(store_metadata.get("headless", True))
        result_queue.put((
            "result",
            {
                "extraction": extract_shopee_menu(store_metadata, output_dir, headless=headless),
                "store_metadata": store_metadata,
            },
        ))
    except BaseException as exc:
        result_queue.put(("error", f"{type(exc).__name__}: {exc}"))
