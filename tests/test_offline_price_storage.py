import os
import sys
import uuid
import openpyxl
from pathlib import Path

# Add project root to sys.path
BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))

from menu_core.database import (
    SessionLocal,
    init_db,
    ItemOfflinePrice,
    normalize_offline_item_name,
    get_offline_prices_by_store,
    bulk_upsert_offline_prices,
)
from menu_core.c5_combiner import _restore_offline_prices, OFFLINE_PRICE_COLUMNS

def test_offline_price_storage_and_restore():
    print("=== TEST 1: Database Initialization & Table Creation ===")
    init_db()
    db = SessionLocal()
    try:
        # Check table exists by querying
        count = db.query(ItemOfflinePrice).count()
        print(f"✅ Table item_offline_prices exists. Existing rows: {count}")

        print("\n=== TEST 2: Normalization ===")
        assert normalize_offline_item_name("  Ayam   Goreng Special! ") == "ayam goreng special"
        assert normalize_offline_item_name(None) == ""
        print("✅ normalize_offline_item_name works as expected.")

        print("\n=== TEST 3: Bulk Upsert & Retrieval ===")
        test_sid = f"TEST_SID_{uuid.uuid4().hex[:8]}"
        test_iid_1 = f"IID_{uuid.uuid4().hex[:6]}"
        test_iid_2 = f"IID_{uuid.uuid4().hex[:6]}"

        records = [
            {
                "store_id": test_sid,
                "item_id": test_iid_1,
                "item_name": "Paket Nasi Ayam",
                "offline_price": 25000,
                "offline_adjustment": 1000,
                "source": "unit_test"
            },
            {
                "store_id": test_sid,
                "item_id": test_iid_2,
                "item_name": "Es Teh Manis",
                "offline_price": 5000,
                "offline_adjustment": 0,
                "source": "unit_test"
            },
            {
                "store_id": test_sid,
                "item_id": None,
                "item_name": "Menu Baru Tanpa ID",
                "offline_price": 18000,
                "offline_adjustment": 500,
                "source": "unit_test"
            }
        ]

        upserted = bulk_upsert_offline_prices(db, records)
        assert upserted == 3, f"Expected 3 upserted, got {upserted}"
        print(f"✅ Upserted {upserted} records.")

        # Test query
        res = get_offline_prices_by_store(db, [test_sid])
        assert (test_sid, test_iid_1) in res['by_id']
        assert res['by_id'][(test_sid, test_iid_1)]['offline_price'] == 25000
        assert res['by_id'][(test_sid, test_iid_1)]['offline_adjustment'] == 1000

        assert (test_sid, test_iid_2) in res['by_id']
        assert res['by_id'][(test_sid, test_iid_2)]['offline_price'] == 5000

        # By name fallback for item without ID
        norm_name_new = normalize_offline_item_name("Menu Baru Tanpa ID")
        assert (test_sid, norm_name_new) in res['by_name']
        assert res['by_name'][(test_sid, norm_name_new)]['offline_price'] == 18000
        print("✅ get_offline_prices_by_store returned correct by_id and by_name records.")

        # Test update existing
        update_records = [
            {
                "store_id": test_sid,
                "item_id": test_iid_1,
                "item_name": "Paket Nasi Ayam (Updated)",
                "offline_price": 27000,
                "offline_adjustment": 1500,
                "source": "unit_test_update"
            }
        ]
        upserted_update = bulk_upsert_offline_prices(db, update_records)
        assert upserted_update == 1
        res_updated = get_offline_prices_by_store(db, [test_sid])
        assert res_updated['by_id'][(test_sid, test_iid_1)]['offline_price'] == 27000
        assert res_updated['by_id'][(test_sid, test_iid_1)]['offline_adjustment'] == 1500
        print("✅ Update existing record in database succeeded.")

        print("\n=== TEST 4: C5 Restore from Database ===")
        test_excel = BASE_DIR / "test_artifacts" / "test_offline_restore.xlsx"
        test_excel.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Item"
        ws.append(["SID", "Item ID", "Item", "Offline Price (Rp)", "Offline Adjustment (Rp)"])
        ws.append([test_sid, test_iid_1, "Paket Nasi Ayam", None, None])
        ws.append([test_sid, "", "Menu Baru Tanpa ID", None, None])
        wb.save(test_excel)
        wb.close()

        restored = _restore_offline_prices(str(test_excel), previous_workbook=None, db=db)
        print(f"Restored count from DB: {restored}")
        assert restored >= 2, f"Expected at least 2 restored, got {restored}"

        # Verify Excel content
        wb_check = openpyxl.load_workbook(test_excel)
        ws_check = wb_check["Item"]
        assert ws_check.cell(row=2, column=4).value == 27000, f"Expected 27000, got {ws_check.cell(row=2, column=4).value}"
        assert ws_check.cell(row=2, column=5).value == 1500, f"Expected 1500, got {ws_check.cell(row=2, column=5).value}"
        assert ws_check.cell(row=3, column=4).value == 18000, f"Expected 18000, got {ws_check.cell(row=3, column=4).value}"
        wb_check.close()
        print("✅ Restored values written correctly into Excel columns.")

        print("\n=== TEST 5: Auto-seed from Backup Workbook into Database ===")
        # Create an older workbook with a brand new item not in DB
        backup_excel = BASE_DIR / "test_artifacts" / "test_backup.xlsx"
        wb_bk = openpyxl.Workbook()
        ws_bk = wb_bk.active
        ws_bk.title = "Item"
        ws_bk.append(["SID", "Item ID", "Item", "Offline Price (Rp)", "Offline Adjustment (Rp)"])
        test_iid_seed = f"IID_SEED_{uuid.uuid4().hex[:6]}"
        ws_bk.append([test_sid, test_iid_seed, "Kopi Susu Gula Aren", 15000, 0])
        wb_bk.save(backup_excel)
        wb_bk.close()

        # Create target excel needing that seed item
        target_excel = BASE_DIR / "test_artifacts" / "test_target.xlsx"
        wb_tg = openpyxl.Workbook()
        ws_tg = wb_tg.active
        ws_tg.title = "Item"
        ws_tg.append(["SID", "Item ID", "Item", "Offline Price (Rp)", "Offline Adjustment (Rp)"])
        ws_tg.append([test_sid, test_iid_seed, "Kopi Susu Gula Aren", None, None])
        wb_tg.save(target_excel)
        wb_tg.close()

        # Restore passing backup_excel
        restored_seed = _restore_offline_prices(str(target_excel), previous_workbook=str(backup_excel), db=db)
        assert restored_seed >= 1, f"Expected at least 1 restored from seed, got {restored_seed}"

        # Check DB was auto-seeded!
        res_seeded = get_offline_prices_by_store(db, [test_sid])
        assert (test_sid, test_iid_seed) in res_seeded['by_id'], "Auto-seed failed to insert into DB!"
        assert res_seeded['by_id'][(test_sid, test_iid_seed)]['offline_price'] == 15000
        print("✅ Auto-seeding from backup workbook to database verified successfully.")

        # Cleanup test records
        db.query(ItemOfflinePrice).filter(ItemOfflinePrice.store_id == test_sid).delete()
        db.commit()
        print("✅ Test cleanup complete.")

        # Cleanup test files
        for p in [test_excel, backup_excel, target_excel]:
            if p.exists():
                p.unlink()

        print("\n🎉 ALL TESTS PASSED!")
    finally:
        db.close()

if __name__ == "__main__":
    test_offline_price_storage_and_restore()
