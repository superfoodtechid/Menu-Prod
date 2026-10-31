import os
import tempfile
import unittest
import uuid
from datetime import datetime, timedelta
from unittest.mock import patch

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import main
from menu_core import database
from menu_core.database import Base, Job
from menu_core.gdrive_service import _atomic_save_workbook


class TestBatchCombine(unittest.TestCase):

    def setUp(self):
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
        Base.metadata.create_all(engine)
        self.Session = sessionmaker(bind=engine)
        self.session_patch = patch.object(database, "SessionLocal", self.Session)
        self.session_patch.start()
        self.batch_id = str(uuid.uuid4())

    def tearDown(self):
        self.session_patch.stop()

    def _add_job(self, status, batch_id=None):
        with self.Session() as db:
            job = Job(
                job_type="PULL",
                platform="grab",
                status=status,
                payload={"batch_id": batch_id or self.batch_id, "batch_label": "Outlet A"},
            )
            db.add(job)
            db.commit()
            return str(job.id)

    def _record(self):
        with self.Session() as db:
            return main._get_batch_combine_record(main._query_batch_jobs(db, self.batch_id))

    @patch("main._execute_combine_and_upload")
    def test_waits_until_every_job_is_terminal(self, mock_exec):
        self._add_job("SUCCESS")
        self._add_job("RUNNING")
        main.run_batch_combine(self.batch_id)
        mock_exec.assert_not_called()

    @patch("main._execute_combine_and_upload")
    def test_combines_success_jobs_once(self, mock_exec):
        mock_exec.return_value = {"ok": True, "combine_id": "c1"}
        ok_id = self._add_job("SUCCESS")
        self._add_job("FAILED")
        self._add_job("SUCCESS", batch_id="other-batch")

        main.run_batch_combine(self.batch_id)
        args, kwargs = mock_exec.call_args
        self.assertEqual(args[0], [ok_id])
        self.assertEqual(args[1], "Outlet A")
        self.assertEqual(kwargs["batch_id"], self.batch_id)

        # Simulate _execute_combine_and_upload having stored its result.
        main._update_batch_combine(self.batch_id, {"status": "UPLOADING"})
        main.run_batch_combine(self.batch_id)
        self.assertEqual(mock_exec.call_count, 1)

        main.run_batch_combine(self.batch_id, force=True)
        self.assertEqual(mock_exec.call_count, 2)

    @patch("main._execute_combine_and_upload")
    def test_fresh_combining_marker_blocks_second_runner(self, mock_exec):
        self._add_job("SUCCESS")
        main._update_batch_combine(self.batch_id, {
            "status": "COMBINING",
            "started_at": datetime.utcnow().isoformat(),
        })
        main.run_batch_combine(self.batch_id, force=True)
        mock_exec.assert_not_called()

        main._update_batch_combine(self.batch_id, {
            "started_at": (datetime.utcnow() - timedelta(hours=1)).isoformat(),
        })
        main.run_batch_combine(self.batch_id)
        mock_exec.assert_called_once()

    @patch("main._execute_combine_and_upload")
    def test_rerun_job_triggers_new_combine(self, mock_exec):
        self._add_job("SUCCESS")
        self._add_job("FAILED")
        main.run_batch_combine(self.batch_id)
        main._update_batch_combine(self.batch_id, {"status": "SUCCESS"})

        self._add_job("SUCCESS")
        main.run_batch_combine(self.batch_id)
        self.assertEqual(mock_exec.call_count, 2)
        self.assertEqual(len(mock_exec.call_args[0][0]), 2)

    @patch("main._execute_combine_and_upload", side_effect=main.HTTPException(status_code=400, detail="no files"))
    def test_combine_error_is_persisted(self, _):
        self._add_job("SUCCESS")
        with self.assertRaises(main.HTTPException):
            main.run_batch_combine(self.batch_id)
        record = self._record()
        self.assertEqual(record["status"], "COMBINE_FAILED")
        self.assertEqual(record["error"], "no files")


class TestAtomicSave(unittest.TestCase):

    def test_replaces_file_and_leaves_no_temp(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "c5.xlsx")
            wb = openpyxl.Workbook()
            wb.active["A1"] = "old"
            wb.save(path)

            wb = openpyxl.load_workbook(path)
            wb.active["A1"] = "new"
            _atomic_save_workbook(wb, path)

            self.assertEqual(openpyxl.load_workbook(path).active["A1"].value, "new")
            self.assertEqual(os.listdir(tmp), ["c5.xlsx"])


if __name__ == "__main__":
    unittest.main()
