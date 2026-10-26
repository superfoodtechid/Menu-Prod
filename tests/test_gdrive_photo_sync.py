import os
import tempfile
import unittest
from unittest.mock import MagicMock, patch
import openpyxl

from menu_core.gdrive_service import (
    _sanitize_name,
    format_photo_folder_label,
    is_photo_sync_enabled,
    calculate_photo_sync_eta,
    get_or_create_subfolder,
    resolve_menu_foto_folder,
    _download_photo_content_base64,
    _restore_cdn_urls_in_excel,
    sync_menu_photos_via_appscript,
    sync_menu_photos_to_drive,
    prepare_and_start_photo_sync,
    update_excel_folder_link,
)
from menu_core.c5_template_helper import (
    upsert_folder_link_sheet,
    read_folder_link_entries,
    FOLDER_LINK_SHEET_NAME,
    FOLDER_LINK_LABEL_HEADER,
    FOLDER_LINK_URL_HEADER,
)
from menu_core.c5_combiner import combine_c5


class TestGDrivePhotoSync(unittest.TestCase):

    def test_sanitize_name(self):
        self.assertEqual(_sanitize_name("Ayam / Goreng: Enak*?"), "Ayam Goreng Enak")
        self.assertEqual(_sanitize_name(""), "Unknown")
        self.assertEqual(_sanitize_name(None), "Unknown")
        self.assertEqual(_sanitize_name("  Burger   King  "), "Burger King")

    @patch("menu_core.gdrive_service.dotenv_values")
    def test_is_photo_sync_enabled_toggle(self, mock_dotenv):
        mock_dotenv.return_value = {"ENABLE_GDRIVE_PHOTO_SYNC": "false"}
        with patch.dict(os.environ, {}, clear=True):
            self.assertFalse(is_photo_sync_enabled())

        mock_dotenv.return_value = {"ENABLE_GDRIVE_PHOTO_SYNC": "true"}
        with patch.dict(os.environ, {}, clear=True):
            self.assertTrue(is_photo_sync_enabled())

    def test_calculate_photo_sync_eta(self):
        eta = calculate_photo_sync_eta(59, batch_size=25)
        self.assertEqual(eta["total_photos"], 59)
        self.assertTrue(eta["estimated_completion"].endswith("WIB"))
        self.assertGreaterEqual(eta["estimated_seconds"], 20)

    @patch("menu_core.gdrive_service.dotenv_values")
    def test_resolve_menu_foto_folder_hierarchy(self, mock_dotenv):
        mock_dotenv.return_value = {"GDRIVE_FOLDER_ID": "root123"}
        mock_service = MagicMock()

        mock_service.files().list().execute.side_effect = [
            {"files": [{"id": "owner_folder_id", "name": "PT Rasa"}]},
            {"files": [{"id": "menu_foto_folder_id", "name": "Menu Foto"}]},
            {"files": [{"id": "outlet_folder_id", "name": "ShopeeFood | Ayam Bakar | 12345 | 2026-10-22 10:00"}]},
        ]

        folder_id, folder_url = resolve_menu_foto_folder(
            service=mock_service,
            owner_name="PT Rasa",
            platform="ShopeeFood",
            listing_name="Ayam Bakar",
            sid="12345",
            timestamp_str="2026-10-22 10:00"
        )

        self.assertEqual(folder_id, "outlet_folder_id")
        self.assertEqual(folder_url, "https://drive.google.com/drive/folders/outlet_folder_id")

    @patch("menu_core.gdrive_service.requests.get")
    def test_download_photo_content_base64(self, mock_requests_get):
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.content = b"fake-image-bytes"
        mock_resp.headers = {"Content-Type": "image/jpeg"}
        mock_requests_get.return_value = mock_resp

        res = _download_photo_content_base64("Ayam Bakar", "https://cf.shopee.co.id/image.jpg", "dish_1")
        self.assertIsNotNone(res)
        self.assertEqual(res["fileName"], "Ayam Bakar.jpg")
        self.assertIn("fileContent", res)

    @patch.dict(os.environ, {"GOOGLE_DRIVE_API_KEY": ""}, clear=False)
    @patch("menu_core.gdrive_service.dotenv_values")
    @patch("menu_core.gdrive_service.requests.post")
    @patch("menu_core.gdrive_service.requests.get")
    def test_sync_menu_photos_via_appscript(self, mock_get, mock_post, mock_dotenv):
        mock_dotenv.return_value = {
            "GDRIVE_APPSCRIPT_URL": "https://script.google.com/macros/s/test/exec",
            "ENABLE_GDRIVE_PHOTO_SYNC": "true",
        }
        mock_post_resp = MagicMock()
        mock_post_resp.status_code = 200
        mock_post_resp.json.return_value = {
            "status": "success",
            "mode": "direct_url_v2",
            "folderUrl": "https://drive.google.com/drive/folders/folder_appscript_123",
            "subFolderId": "folder_appscript_123",
            "uploadedCount": 1,
            "failedFiles": [],
        }
        mock_post.return_value = mock_post_resp

        folder_url = sync_menu_photos_via_appscript(
            owner_name="PT Rasa",
            platform="ShopeeFood",
            listing_name="Ayam Bakar",
            sid="12345",
            dishes=[{"item_name": "Ayam Bakar", "photo_url": "https://example.com/img.jpg", "item_id": "1"}],
            timestamp_str="2026-10-22 10:00",
            run_async=False,
        )
        self.assertEqual(folder_url, "https://drive.google.com/drive/folders/folder_appscript_123")
        mock_get.assert_not_called()

    @patch.dict(os.environ, {"GOOGLE_DRIVE_API_KEY": ""}, clear=False)
    @patch("menu_core.gdrive_service.dotenv_values")
    @patch("menu_core.gdrive_service.requests.post")
    @patch("menu_core.gdrive_service.requests.get")
    def test_direct_url_per_item_base64_fallback(self, mock_get, mock_post, mock_dotenv):
        mock_dotenv.return_value = {
            "GDRIVE_APPSCRIPT_URL": "https://script.google.com/macros/s/test/exec",
            "ENABLE_GDRIVE_PHOTO_SYNC": "true",
        }
        mock_get_resp = MagicMock()
        mock_get_resp.status_code = 200
        mock_get_resp.content = b"fallback-image-bytes"
        mock_get_resp.headers = {"Content-Type": "image/jpeg"}
        mock_get.return_value = mock_get_resp

        resp_init = MagicMock(status_code=200)
        resp_init.json.return_value = {
            "status": "success",
            "mode": "direct_url_v2",
            "folderUrl": "https://drive.google.com/drive/folders/f1",
            "subFolderId": "f1",
            "uploadedCount": 0,
        }
        resp_batch_url = MagicMock(status_code=200)
        resp_batch_url.json.return_value = {
            "status": "success",
            "mode": "direct_url_v2",
            "uploadedCount": 1,
            "failedFiles": [{"fileName": "Es Teh.jpg", "fileUrl": "https://example.com/esteh.jpg", "itemId": "2"}],
        }
        resp_fallback_b64 = MagicMock(status_code=200)
        resp_fallback_b64.json.return_value = {
            "status": "success",
            "mode": "direct_url_v2",
            "uploadedCount": 1,
        }
        mock_post.side_effect = [resp_init, resp_batch_url, resp_fallback_b64]

        folder_url = sync_menu_photos_via_appscript(
            owner_name="PT Rasa",
            platform="GrabFood",
            listing_name="Outlet A",
            sid="S1",
            dishes=[
                {"item_name": "Ayam Bakar", "photo_url": "https://example.com/ayam.jpg", "item_id": "1"},
                {"item_name": "Es Teh", "photo_url": "https://example.com/esteh.jpg", "item_id": "2"},
            ],
            run_async=False,
        )
        self.assertEqual(folder_url, "https://drive.google.com/drive/folders/f1")
        self.assertEqual(mock_post.call_count, 3)
        mock_get.assert_called_once()

    def test_restore_cdn_urls_in_excel_on_total_failure(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Item"
            ws.append(["SID", "Item ID", "Item", "Photo Link"])
            ws.append(["", "", "", ""])
            ws.append(["", "", "", ""])
            ws.append(["S1", "101", "Nasi Goreng", "https://drive.google.com/drive/folders/empty_folder"])
            wb.save(tmp_path)
            wb.close()

            _restore_cdn_urls_in_excel(
                tmp_path,
                [{"item_id": "101", "item_name": "Nasi Goreng", "photo_url": "https://cdn.grab.com/nasigoreng.jpg"}],
            )

            wb_check = openpyxl.load_workbook(tmp_path)
            ws_check = wb_check["Item"]
            self.assertEqual(ws_check.cell(row=4, column=4).value, "https://cdn.grab.com/nasigoreng.jpg")
            wb_check.close()
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    @patch("menu_core.gdrive_service.is_photo_sync_enabled", return_value=True)
    @patch("menu_core.gdrive_service.get_drive_service")
    @patch("menu_core.gdrive_service.sync_menu_photos_via_appscript")
    def test_sync_menu_photos_to_drive_uses_appscript_fallback(self, mock_appscript, mock_sa, _mock_enabled):
        mock_sa.return_value = None
        mock_appscript.return_value = "https://drive.google.com/drive/folders/fallback_url"

        res = sync_menu_photos_to_drive(
            owner_name="Owner",
            platform="ShopeeFood",
            listing_name="Outlet",
            sid="123",
            dishes=[]
        )
        self.assertEqual(res, "https://drive.google.com/drive/folders/fallback_url")
        mock_appscript.assert_called_once()

    @patch("menu_core.c5_combiner.get_c5_template_path")
    def test_folder_link_sheet_and_combine_c5_aggregation(self, mock_template_path):
        mock_template_path.return_value = "/tmp/non_existent_template.xlsx"
        with tempfile.TemporaryDirectory() as tmpdir:
            go_path = os.path.join(tmpdir, "O.C5 GoFood.xlsx")
            shopee_path = os.path.join(tmpdir, "O.C5 ShopeeFood.xlsx")
            combined_path = os.path.join(tmpdir, "Combined.xlsx")

            wb_go = openpyxl.Workbook()
            ws_go_item = wb_go.active
            ws_go_item.title = "Item"
            ws_go_item.append(["OFD", "Listing Name", "SID", "Category ID", "Category", "Item ID", "Item", "Photo Link"])
            ws_go_item.append(["GoFood", "Mie Gacoan", "G100", "C1", "Mie", "I1", "Mie Suit", "https://cdn.gofood.co.id/miesuit.jpg"])
            ws_go_mod = wb_go.create_sheet("Modifier")
            ws_go_mod.append(["OFD", "Listing Name", "SID", "Item", "Modifier Group ID", "Modifier Group", "Modifier ID", "Modifier"])
            go_label = format_photo_folder_label("GoFood", "Mie Gacoan", "G100", "2026-10-23 09:00")
            upsert_folder_link_sheet(wb_go, [(go_label, "https://drive.google.com/drive/folders/go_folder_1")])
            wb_go.save(go_path)
            wb_go.close()

            wb_sh = openpyxl.Workbook()
            ws_sh_item = wb_sh.active
            ws_sh_item.title = "Item"
            ws_sh_item.append(["OFD", "Listing Name", "SID", "Category ID", "Category", "Item ID", "Item", "Photo Link"])
            ws_sh_item.append(["ShopeeFood", "Mie Gacoan", "S200", "C2", "Mie", "I2", "Mie Hompimpa", "https://cdn.shopee.co.id/hompimpa.jpg"])
            ws_sh_mod = wb_sh.create_sheet("Modifier")
            ws_sh_mod.append(["OFD", "Listing Name", "SID", "Item", "Modifier Group ID", "Modifier Group", "Modifier ID", "Modifier"])
            sh_label = format_photo_folder_label("ShopeeFood", "Mie Gacoan", "S200", "2026-10-23 09:01")
            upsert_folder_link_sheet(wb_sh, [(sh_label, "https://drive.google.com/drive/folders/sh_folder_2")])
            wb_sh.save(shopee_path)
            wb_sh.close()

            ok = combine_c5([shopee_path, go_path], combined_path)
            self.assertTrue(ok)

            entries = read_folder_link_entries(combined_path)
            self.assertEqual(
                entries,
                [
                    (go_label, "https://drive.google.com/drive/folders/go_folder_1"),
                    (sh_label, "https://drive.google.com/drive/folders/sh_folder_2"),
                ],
            )

            wb_comb = openpyxl.load_workbook(combined_path, read_only=True)
            self.assertIn(FOLDER_LINK_SHEET_NAME, wb_comb.sheetnames)
            ws_fl = wb_comb[FOLDER_LINK_SHEET_NAME]
            header_row = next(ws_fl.iter_rows(min_row=1, max_row=1, values_only=True))
            self.assertEqual(header_row[0], FOLDER_LINK_LABEL_HEADER)
            self.assertEqual(header_row[1], FOLDER_LINK_URL_HEADER)
            wb_comb.close()

    def test_update_excel_folder_link(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Item"
            ws.append(["Item ID", "Item", "Photo Link"])
            wb.save(tmp_path)
            wb.close()

            ok = update_excel_folder_link(
                tmp_path,
                "ShopeeFood | Bebek | 123 | 2026-10-26 07:00",
                "https://drive.google.com/drive/folders/folder_123"
            )
            self.assertTrue(ok)

            entries = read_folder_link_entries(tmp_path)
            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0][0], "ShopeeFood | Bebek | 123 | 2026-10-26 07:00")
            self.assertEqual(entries[0][1], "https://drive.google.com/drive/folders/folder_123")
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    @patch("menu_core.gdrive_service.is_photo_sync_enabled", return_value=True)
    @patch("menu_core.gdrive_service.sync_menu_photos_to_drive")
    def test_prepare_and_start_photo_sync_async_decoupling(self, mock_sync_drive, _mock_enabled):
        mock_sync_drive.return_value = "https://drive.google.com/drive/folders/async_url"

        folder_url, meta = prepare_and_start_photo_sync(
            owner_name="Test Owner",
            platform="ShopeeFood",
            listing_name="Test Outlet",
            sid="1001",
            dishes=[{"item_name": "Ayam", "photo_url": "https://cdn.shopee.co.id/ayam.jpg", "item_id": "1"}],
            run_async=True,
        )

        # Non-blocking returns immediately with status PROCESSING
        self.assertIsNone(folder_url)
        self.assertIsNotNone(meta)
        self.assertEqual(meta["status"], "PROCESSING")
        self.assertEqual(meta["total_photos"], 1)
        self.assertIsNone(meta["drive_folder_url"])
        self.assertTrue(meta["target_folder_name"].startswith("ShopeeFood | Test Outlet | 1001 |"))


if __name__ == "__main__":
    unittest.main()
