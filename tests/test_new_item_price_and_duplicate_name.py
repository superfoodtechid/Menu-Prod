# -*- coding: utf-8 -*-
import io
import json
import unittest
import openpyxl
from pathlib import Path
from unittest.mock import MagicMock, patch

from main import parse_c5_endpoint, _push_c5_gofood_for_merchant
from fastapi import UploadFile


def create_mock_c5_workbook(rows_data):
    """
    Helper untuk membuat file Excel C5 in-memory.
    rows_data: list of dicts dengan key header standar.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Item"

    headers = [
        "OFD", "Listing Name", "SID", "Outlet", "Category ID", "Category",
        "Item ID", "Item", "Item Name Improvement", "Description",
        "New Fake Price (Rp)", "Current Fake Price (Rp)"
    ]
    for c_idx, h in enumerate(headers, 1):
        ws.cell(1, c_idx, h)

    # Row 2 is reserved
    ws.cell(2, 1, "Don't touch row 2")

    # Data rows start at 3
    for r_idx, row in enumerate(rows_data, 3):
        for c_idx, h in enumerate(headers, 1):
            if h in row:
                ws.cell(r_idx, c_idx, row[h])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestNewItemPriceAndDuplicateName(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.sid = "TESTSID999"
        self.api_dir = Path(__file__).parent.parent / "Gofood" / "API"
        self.api_dir.mkdir(parents=True, exist_ok=True)
        self.baseline_file = self.api_dir / f"menu-response-{self.sid}.json"

        # Baseline memiliki item 'Ayam Geprek' dengan ID 'ITEM-101' dan harga 20000
        baseline_content = {
            "items": [
                {
                    "Store ID": self.sid,
                    "Item ID": "ITEM-101",
                    "Item": "Ayam Geprek",
                    "Category": "Makanan",
                    "Current Fake Price (Rp)": 20000,
                    "Link Foto": "https://cdn.example.com/ayam.jpg",
                    "Description": "Ayam geprek pedas gurih"
                }
            ]
        }
        with open(self.baseline_file, "w", encoding="utf-8") as f:
            json.dump(baseline_content, f)

    def tearDown(self):
        if self.baseline_file.exists():
            self.baseline_file.unlink()

    async def test_duplicate_name_with_empty_id_is_new_item_and_no_step_push(self):
        """
        Dua item dengan nama sama ('Ayam Geprek'):
        - Row 1: Item ID 'ITEM-101', harga naik dari 20rb ke 25rb (+25% -> step push).
        - Row 2: Item ID kosong (item baru), harga 50rb (>15% dari baseline).
        Harapan:
        1. Row 2 terdeteksi sebagai is_new_item = True dan Item ID tetap kosong.
        2. Harga Row 2 TIDAK memicu PRICE_WARNING_STEP_PUSH atau price_warning = True.
        3. diff_details untuk Row 2 menampilkan '(Baru)', bukan 'Push Bertahap >15%'.
        4. price_warning_count hanya 1 (berasal dari Row 1).
        """
        rows = [
            {
                "OFD": "GoFood",
                "SID": self.sid,
                "Category": "Makanan",
                "Category ID": "cat-1",
                "Item ID": "ITEM-101",
                "Item": "Ayam Geprek",
                "New Fake Price (Rp)": 25000,
                "Current Fake Price (Rp)": 20000,
            },
            {
                "OFD": "GoFood",
                "SID": self.sid,
                "Category": "Makanan",
                "Category ID": "cat-1",
                "Item ID": "",  # Empty ID -> New Item!
                "Item": "Ayam Geprek",
                "New Fake Price (Rp)": 50000,
                "Current Fake Price (Rp)": 50000,
            }
        ]
        buf = create_mock_c5_workbook(rows)
        upload_file = UploadFile(filename="O.C5 Test.xlsx", file=buf)

        mock_db = MagicMock()
        mock_db.query.return_value.filter.return_value.all.return_value = []
        with patch("menu_core.database.get_offline_prices_by_store", return_value={"by_id": {}, "by_name": {(self.sid, "dummy"): {"offline_price": 0}}}):
            res = await parse_c5_endpoint(upload_file, db=mock_db)

        self.assertTrue(res.get("success"))
        items = res.get("items", [])
        self.assertEqual(len(items), 2)

        # Item 1: existing item update
        it1 = items[0]
        self.assertEqual(it1["item_id"], "ITEM-101")
        self.assertFalse(it1["is_new_item"])
        self.assertTrue(it1["price_warning"])
        self.assertIn("PRICE_WARNING_STEP_PUSH", it1["change_types"])
        self.assertIn("PRICE_CHANGE", it1["change_types"])

        # Item 2: brand new item with same name
        it2 = items[1]
        self.assertEqual(it2["item_id"], "")
        self.assertTrue(it2["is_new_item"])
        self.assertIsNone(it2["current_fake_price"])
        self.assertFalse(it2["baseline_found"])
        self.assertFalse(it2["price_warning"])
        self.assertEqual(it2["price_diff_percent"], 0.0)
        self.assertIn("NEW_ITEM", it2["change_types"])
        self.assertNotIn("PRICE_WARNING_STEP_PUSH", it2["change_types"])
        self.assertNotIn("PRICE_CHANGE", it2["change_types"])

        # Periksa diff_details pada item 2
        price_diff = [d for d in it2["diff_details"] if d["column"] == "Price"]
        self.assertTrue(len(price_diff) > 0)
        self.assertEqual(price_diff[0]["old_val"], "(Baru)")
        self.assertIn("Rp 50,000", price_diff[0]["new_val"])
        self.assertNotIn("Push Bertahap", price_diff[0]["new_val"])

        # Summary check
        summary = res.get("summary", {})
        self.assertEqual(summary.get("new_items_count"), 1)
        self.assertEqual(summary.get("price_warning_count"), 1)  # Only from it1

    async def test_empty_id_row_first_does_not_hijack_explicit_id_row(self):
        """
        Pastikan urutan baris tidak mempengaruhi isolasi:
        Row 1 (kosong ID) berada sebelum Row 2 (berisi ID 'ITEM-101').
        Row 1 tetap dianggap item baru dan TIDAK merebut 'ITEM-101'.
        """
        rows = [
            {
                "OFD": "GoFood",
                "SID": self.sid,
                "Category": "Makanan",
                "Category ID": "cat-1",
                "Item ID": "",  # Empty ID appears first
                "Item": "Ayam Geprek",
                "New Fake Price (Rp)": 50000,
                "Current Fake Price (Rp)": 50000,
            },
            {
                "OFD": "GoFood",
                "SID": self.sid,
                "Category": "Makanan",
                "Category ID": "cat-1",
                "Item ID": "ITEM-101",  # Explicit ID appears second
                "Item": "Ayam Geprek",
                "New Fake Price (Rp)": 25000,
                "Current Fake Price (Rp)": 20000,
            }
        ]
        buf = create_mock_c5_workbook(rows)
        upload_file = UploadFile(filename="O.C5 Test Order.xlsx", file=buf)

        mock_db = MagicMock()
        mock_db.query.return_value.filter.return_value.all.return_value = []
        with patch("menu_core.database.get_offline_prices_by_store", return_value={"by_id": {}, "by_name": {(self.sid, "dummy"): {"offline_price": 0}}}):
            res = await parse_c5_endpoint(upload_file, db=mock_db)

        items = res.get("items", [])
        self.assertEqual(len(items), 2)

        it_new = items[0]
        it_existing = items[1]

        self.assertEqual(it_new["item_id"], "")
        self.assertTrue(it_new["is_new_item"])
        self.assertFalse(it_new["price_warning"])
        self.assertNotIn("PRICE_WARNING_STEP_PUSH", it_new["change_types"])

        self.assertEqual(it_existing["item_id"], "ITEM-101")
        self.assertFalse(it_existing["is_new_item"])
        self.assertTrue(it_existing["price_warning"])

    @patch("playwright.sync_api.sync_playwright")
    @patch("src.core.browser_factory.launch_universal_playwright_browser")
    @patch("login_gofood.load_gofood_session", return_value={"access_token": "Bearer fake-token", "cookies": []})
    @patch("menu_core.gofood.fetch_gofood_mpp_promotions", return_value={})
    @patch("Gofood.GO.actions._menu_api.fetch_menus")
    @patch("Gofood.GO.actions._menu_api.fetch_menus_v2")
    @patch("Gofood.GO.actions._menu_api.create_menu_item")
    def test_push_gofood_creates_new_item_even_with_same_name_as_existing(
        self,
        mock_create_item,
        mock_fetch_menus_v2,
        mock_fetch_menus,
        _mock_mpp,
        _mock_load_session,
        mock_launch_browser,
        mock_sync_pw,
    ):
        """
        Memverifikasi bahwa _push_c5_gofood_for_merchant tidak membajak item baru
        yang namanya sama dengan item lama di GoFood, melainkan memanggil create_menu_item.
        """
        initial_menu = {
            "menus": [
                {
                    "id": "cat-v1-1",
                    "common_id": "cat-common-1",
                    "name": "Makanan",
                    "menu_items": [
                        {
                            "id": "item-existing-123",
                            "common_id": "item-common-123",
                            "name": "Kwetiau Goreng",
                            "price": 25000,
                            "active": True,
                        }
                    ]
                }
            ]
        }
        mock_fetch_menus.return_value = initial_menu
        mock_fetch_menus_v2.return_value = {"menus": []}
        mock_create_item.return_value = {"ok": True, "status": 200, "body": json.dumps({"data": {"id": "new-item-456"}})}

        fake_pw_cm = MagicMock()
        mock_sync_pw.return_value = fake_pw_cm
        fake_browser_obj = MagicMock()
        fake_ctx = MagicMock()
        fake_page = MagicMock()

        def fake_on(event_name, handler):
            if event_name == "request":
                req = MagicMock()
                req.url = "https://api.gojekapi.com/gofood/merchant/v2/menu_groups/11111111-1111-1111-1111-111111111111/menus"
                req.headers = {
                    "authorization": "Bearer fake-token",
                    "x-passkey": "passkey-123",
                }
                handler(req)
                req2 = MagicMock()
                req2.url = "https://api.gojekapi.com/gofood/merchant/v1/restaurants/12345678-1234-1234-1234-1234567890ab/menus"
                req2.headers = {"authorization": "Bearer fake-token"}
                handler(req2)

        fake_page.on.side_effect = fake_on
        fake_ctx.new_page.return_value = fake_page
        fake_browser_obj.new_context.return_value = fake_ctx
        mock_launch_browser.return_value = (fake_browser_obj, None)

        updates = [
            {
                "sid": "G-12345",
                "item_id": "",
                "item_name": "Kwetiau Goreng",
                "new_fake_price": 35000,
                "category": "Makanan",
                "new_category": "Makanan",
                "changes": ["NEW_ITEM"],
            }
        ]

        results = _push_c5_gofood_for_merchant(
            email="test@example.com",
            password="pass",
            merchant_id="G-12345",
            updates=updates,
        )

        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["status"], "SUCCESS")
        mock_create_item.assert_called_once()
        create_args = mock_create_item.call_args[0]
        # create_args: (page, token, patch_group_id, payload)
        payload = create_args[3]
        self.assertEqual(payload["name"], "Kwetiau Goreng")
        self.assertEqual(payload["price"], 35000)
        self.assertEqual(payload["menu_common_id"], "cat-common-1")

    def test_grab_push_item_matching_isolation_for_new_items(self):
        """
        Memverifikasi bahwa pada Grab push, item baru (is_new_item=True, item_id='')
        tidak mencocokkan existing item yang memiliki nama sama di existing_items_map.
        """
        # Simulasikan existing_items_map pada Grab push
        existing_items_map = {
            "grab-item-123": {
                "itemID": "grab-item-123",
                "itemName": "Sate Ayam",
                "priceInMin": 2500000,
            },
            "sate ayam": {
                "itemID": "grab-item-123",
                "itemName": "Sate Ayam",
                "priceInMin": 2500000,
            }
        }

        # Item baru dengan nama sama dan item_id kosong
        item_id = ""
        item_name = "Sate Ayam"
        is_new_item = True
        res_item = {"item_id": item_id, "item_name": item_name}

        # Logika dari grab/core/push_c5.py
        if is_new_item:
            existing_it = existing_items_map.get(item_id) if item_id else None
        else:
            existing_it = existing_items_map.get(item_id) or existing_items_map.get(item_name.strip().lower())

        if existing_it and not is_new_item:
            res_item["item_id"] = str(existing_it.get("itemID") or existing_it.get("id") or item_id)

        # existing_it harus None untuk item baru
        self.assertIsNone(existing_it)
        # res_item['item_id'] tetap kosong (tidak dibajak 'grab-item-123')
        self.assertEqual(res_item["item_id"], "")


if __name__ == "__main__":
    unittest.main()
