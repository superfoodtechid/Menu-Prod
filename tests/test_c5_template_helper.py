# -*- coding: utf-8 -*-
import unittest
import os
import shutil
import tempfile
from pathlib import Path
import openpyxl

from menu_core.c5_template_helper import (
    get_c5_template_path,
    prepare_c5_sheet,
    get_c5_headers,
    DATA_START_ROW,
    TEMPLATE_RESERVED_ROW,
    is_c5_reserved_row,
)


class TestC5TemplateHelperFast(unittest.TestCase):
    def test_reserved_row_detection(self):
        """Memverifikasi deteksi baris 2 ('Don't touch row 2')."""
        self.assertTrue(is_c5_reserved_row(["Don't touch row 2", None, None]))
        self.assertTrue(is_c5_reserved_row(["don't touch row 2", "x", "y"]))
        self.assertFalse(is_c5_reserved_row(["GoFood", "Outlet A", "http...", "123"]))
        self.assertFalse(is_c5_reserved_row(["ShopeeFood", "Outlet B", "http...", "456"]))
        self.assertTrue(is_c5_reserved_row([]))
        self.assertTrue(is_c5_reserved_row(None))

    def test_prepare_c5_sheet_unit_logic(self):
        """Memverifikasi bahwa prepare_c5_sheet membersihkan data mulai baris 3 tanpa menyentuh baris 2."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Item"
        # Baris 1: Header
        ws.cell(row=1, column=1, value="OFD")
        ws.cell(row=1, column=2, value="Listing Name")
        # Baris 2: Reserved
        ws.cell(row=2, column=1, value="Don't touch row 2")
        ws.cell(row=2, column=2, value=None)
        # Baris 3, 4, 5: Data menu lama
        ws.cell(row=3, column=1, value="GoFood")
        ws.cell(row=3, column=2, value="Ayam Bakar")
        ws.cell(row=4, column=1, value="GoFood")
        ws.cell(row=4, column=2, value="Ayam Goreng")
        ws.cell(row=5, column=1, value="GoFood")
        ws.cell(row=5, column=2, value="Es Teh")

        self.assertEqual(ws.max_row, 5)

        # Jalankan pembersihan
        prepare_c5_sheet(ws)

        # Baris 1 dan 2 harus tetap utuh
        self.assertEqual(ws.cell(row=1, column=1).value, "OFD")
        self.assertEqual(ws.cell(row=2, column=1).value, "Don't touch row 2")
        # Baris 3 ke atas harus sudah dihapus
        self.assertEqual(ws.max_row, 2)

    def test_get_c5_headers_unit_logic(self):
        """Memverifikasi ekstraksi header string dan mengabaikan non-string/ArrayFormula."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="OFD")
        ws.cell(row=1, column=2, value="Listing Name")
        ws.cell(row=1, column=3, value=12345)  # non-string
        ws.cell(row=1, column=4, value=None)   # empty

        headers = get_c5_headers(ws)
        self.assertEqual(headers, {"OFD": 1, "Listing Name": 2})

    def test_real_template_structure_and_sheets(self):
        """Memverifikasi file template O. C5 .xlsx asli dan keberadaan seluruh 10 sheets."""
        workspace_dir = Path(__file__).resolve().parents[1]
        template_path = get_c5_template_path(workspace_dir)
        self.assertTrue(template_path.exists())
        self.assertEqual(template_path.name, "O. C5 .xlsx")

        wb = openpyxl.load_workbook(str(template_path), read_only=True)
        expected_sheets = [
            'Item', 'Modifier', 'Item ME', 'Modifier ME',
            'Structure', 'Pricing', 'Meaning', 'Tech Hierarchy', 'Guide', 'Method'
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)
        wb.close()


if __name__ == '__main__':
    unittest.main()
